"""Run on Windows: .venv/Scripts/python.exe tests/smoke.py (real model, no mocks)."""

import hashlib
import os
import tempfile
import time
from pathlib import Path

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation
from PySide6.QtWidgets import QApplication

from easything.app import Window
from easything.index import Index, canonical, in_scope, selected_files
from easything.model import Embedder, MODEL_FILE, verify_model
from easything.readers import read_units
from easything.worker import Worker


class CountedEmbedder(Embedder):
    calls = 0

    def encode(self, *args, **kwargs):
        self.calls += 1
        return super().encode(*args, **kwargs)


class UnavailableGPU(Embedder):
    def _load(self, gpu):
        if gpu:
            raise RuntimeError("GPU unavailable (smoke check)")
        return super()._load(False)


def wait_for(app, predicate, timeout=40):
    until = time.monotonic() + timeout
    while time.monotonic() < until:
        app.processEvents()
        if predicate():
            return
        time.sleep(0.05)
    raise AssertionError("Timed out waiting for filesystem indexing")


def run():
    app = QApplication.instance() or QApplication([])
    real_root = Path(os.environ["LOCALAPPDATA"]) / "EasyThing"
    assert verify_model(real_root / "models" / MODEL_FILE)
    with tempfile.TemporaryDirectory(prefix="easything-smoke-") as temporary:
        base = Path(temporary)
        root = base / "data"
        for name in ("models", "index", "cache", "settings", "logs"):
            (root / name).mkdir(parents=True)
        os.link(real_root / "models" / MODEL_FILE, root / "models" / MODEL_FILE)
        fallback = UnavailableGPU(root)
        assert len(fallback.encode("IFC certification", query=True)) == 768
        assert fallback.backend == "CPU"
        fallback.close()
        selected, outside = base / "selected's_%", base / "outside"
        selected.mkdir()
        outside.mkdir()
        (outside / "private.md").write_text("Never index this", encoding="utf-8")
        nested = selected / "nested"
        nested.mkdir()
        (nested / "deep.md").write_text("ESP32-S3 sensor", encoding="utf-8")
        md = selected / "project.md"
        md.write_text("A project proposal for Revit interoperability and MEP coordination, IFC4 certification.", encoding="utf-8")
        (selected / "notes.txt").write_text("RTX 5090 and error 0x80070005", encoding="utf-8")
        deck = selected / "revit.pptx"
        ppt = Presentation()
        for content in ("Revit interoperability and MEP", "IFC certification TS EN ISO 16739-1"):
            slide = ppt.slides.add_slide(ppt.slide_layouts[1])
            slide.shapes.title.text = content
        ppt.save(deck)
        pdf = selected / "ifc.pdf"
        document = pymupdf.open()
        for content in ("Building interoperability", "IFC certification CA742"):
            page = document.new_page()
            page.insert_text((72, 72), content)
        document.save(pdf)
        document.close()
        picture = selected / "red-square.png"
        image = Image.new("RGB", (400, 400), "white")
        ImageDraw.Draw(image).rectangle((70, 70, 330, 330), fill="red")
        image.save(picture)
        for suffix in ("jpg", "jpeg", "webp"):
            image.save(selected / f"red-square.{suffix}")
        from pillow_heif import from_pillow
        from_pillow(image).save(selected / "red-square.heic")
        blue = Image.new("RGB", (400, 400), "white")
        ImageDraw.Draw(blue).ellipse((70, 70, 330, 330), fill="blue")
        blue.save(selected / "photo002.png")
        folders = [{"path": str(selected), "enabled": True, "recursive": True}]
        assert not in_scope(outside / "private.md", folders)
        assert len(list(selected_files(folders))) == 11
        flat = [{**folders[0], "recursive": False}]
        assert canonical(nested / "deep.md") not in list(selected_files(flat))
        hashes = {str(path): hashlib.sha256(path.read_bytes()).digest() for path in selected.rglob("*") if path.is_file()}
        embedder = CountedEmbedder(root, cpu_only=True)
        index = Index(root, embedder)
        for path in selected_files(folders):
            print("Indexing", Path(path).name, flush=True)
            index.index_file(path)
        assert index.table.count_rows() == 13
        assert index.index_file(md) == "unchanged"
        assert len(list(read_units(deck, root / "cache"))) == 2
        assert len(list(read_units(pdf, root / "cache"))) == 2
        print("Search", flush=True)
        hits = index.search("project proposal about Revit interoperability", folders)
        assert canonical(md) in [hit["path"] for hit in hits[:3]]
        assert index.search("IFC4", folders)[0]["path"] == canonical(md)
        assert index.search("0x80070005", folders)[0]["path"] == canonical(selected / "notes.txt")
        assert index.search("a crimson geometric shape on a white background", folders, "Images")[0]["name"] != "photo002.png"
        assert index.search("ESP32-S3", flat) == [] or all(hit["path"] != canonical(nested / "deep.md") for hit in index.search("ESP32-S3", flat))
        assert not index.search("Revit", [{**folders[0], "enabled": False}])
        index.close()
        embedder.close()
        embedder = Embedder(root, cpu_only=True)
        index = Index(root, embedder)
        assert index.search("IFC certification", folders)
        assert all(hashlib.sha256(Path(path).read_bytes()).digest() == digest for path, digest in hashes.items())
        # Real inference, one changed page: the unchanged page must reuse its vector.
        document = pymupdf.open(pdf)
        document[1].insert_text((72, 100), "Updated IFC certification")
        updated = selected / "updated.pdf"
        document.save(updated)
        document.close()
        updated.replace(pdf)
        counted = CountedEmbedder(root, cpu_only=True)
        index.embedder = counted
        index.index_file(pdf)
        assert counted.calls == 1
        counted.close()
        index.close()
        embedder.close()
        print("Restart persistence, page hash reuse and source integrity OK", flush=True)
        worker = Worker(root, {"folders": folders}, cpu_only=True)
        issues = []
        reconciled = []
        worker.progress.connect(lambda done, total, name: reconciled.append((done, total)))
        worker.issue.connect(issues.append)
        worker.start()
        fresh = selected / "fresh.md"
        try:
            time.sleep(1)
            wait_for(app, lambda: bool(reconciled))
            assert reconciled[0] == (0, 0), "Restart must not queue unchanged files"
            fresh.write_text("Fresh local search document", encoding="utf-8")
            def count_path(path):
                import lancedb
                return lancedb.connect(root / "index" / "vectors").open_table("units").search().where(
                    "path = '" + canonical(path).replace("'", "''") + "'"
                ).limit(None).to_list()
            wait_for(app, lambda: bool(count_path(fresh)))
            fresh.write_text("Modified document CA742", encoding="utf-8")
            wait_for(app, lambda: "Modified" in (count_path(fresh) or [{}])[0].get("content", ""))
            moved = selected / "moved.md"
            fresh.rename(moved)
            wait_for(app, lambda: not count_path(fresh) and bool(count_path(moved)))
            moved.unlink()
            wait_for(app, lambda: not count_path(moved))
            (selected / "broken.jpg").write_bytes(b"invalid image")
            wait_for(app, lambda: bool(issues), timeout=20)
            print("Watch create / modify / move / delete and broken file isolation OK", flush=True)
        finally:
            worker.submit("quit")
            assert worker.wait(30000)
        window = Window(root, {"folders": []}, cpu_only=True)
        window.show()
        app.processEvents()
        assert window.isVisible() and window.settings["folders"] == []
        window.settings["folders"].append(folders[0])
        window.update_folders()
        app.processEvents()
        window.exit_app()
        assert window.worker.wait(30000)
        window.tray.hide()
        window.close()
        print("UI startup / selected folder settings OK", flush=True)
    print("PASS: real CPU embeddings, all formats, hybrid search, persistence, watcher, read-only files and UI", flush=True)


if __name__ == "__main__":
    run()
