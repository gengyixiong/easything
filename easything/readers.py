"""Thin, read-only readers. Each page or slide is one search unit."""

import hashlib
import io
import logging
import tempfile
from pathlib import Path

import pymupdf
from PIL import Image, ImageOps
from pillow_heif import register_heif_opener
from pptx import Presentation

register_heif_opener()
IMAGES = {".jpg", ".jpeg", ".png", ".webp", ".heic"}
SUPPORTED = IMAGES | {".pptx", ".pdf", ".txt", ".md"}


def image_bytes(image):
    image = ImageOps.exif_transpose(image).convert("RGB")
    image.thumbnail((1600, 1600))
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


def slide_text(shapes):
    parts = []
    for shape in shapes:
        if shape.has_text_frame:
            parts.append(shape.text)
        if shape.has_table:
            parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
        if shape.shape_type == 6:  # GROUP
            parts.append(slide_text(shape.shapes))
    return "\n".join(parts).strip()


def unit(number, label, text, image=None):
    digest = hashlib.sha256(text.encode("utf-8"))
    if image:
        digest.update(image)
    return {"number": number, "match": label, "content": text,
            "image": image, "content_hash": digest.hexdigest()}


def read_units(path, cache):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in IMAGES:
        with Image.open(path) as image:
            yield unit(1, "Image", "", image_bytes(image))
    elif suffix == ".pdf":
        with pymupdf.open(path) as document:
            if document.needs_pass:
                raise ValueError("This PDF is password protected.")
            for number, page in enumerate(document, 1):
                scale = min(2, 1600 / max(page.rect.width, page.rect.height))
                image = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).tobytes("png")
                yield unit(number, f"Page {number}", page.get_text(), image)
    elif suffix == ".pptx":
        presentation = Presentation(path)
        app = deck = None
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        try:
            try:
                app = win32com.client.DispatchEx("PowerPoint.Application")
                app.AutomationSecurity = 3
                deck = app.Presentations.Open(str(path.resolve()), ReadOnly=True, WithWindow=False)
            except Exception:
                logging.info("PowerPoint rendering unavailable: %s", path, exc_info=True)
            with tempfile.TemporaryDirectory(dir=cache) as tmp:
                height = round(1280 * presentation.slide_height / presentation.slide_width)
                for number, slide in enumerate(presentation.slides, 1):
                    image = None
                    if deck is not None:
                        try:
                            output = Path(tmp) / f"{number}.png"
                            deck.Slides(number).Export(str(output), "PNG", 1280, height)
                            with Image.open(output) as rendered:
                                image = image_bytes(rendered)
                        except Exception:
                            logging.exception("Slide rendering unavailable")
                    item = unit(number, f"Slide {number}", slide_text(slide.shapes), image)
                    item["visual"] = image is not None
                    yield item
        finally:
            if deck is not None:
                try:
                    deck.Close()
                except Exception:
                    logging.exception("Could not close read-only PowerPoint presentation")
            if app is not None:
                try:
                    app.Quit()
                except Exception:
                    logging.exception("Could not close EasyThing PowerPoint instance")
            pythoncom.CoUninitialize()
    elif suffix in {".txt", ".md"}:
        data = path.read_bytes()
        encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        text = data.decode(encoding)
        # UTF-8 bytes bound token count conservatively, including Chinese text.
        chunks, current, length = [], [], 0
        for character in text:
            size = len(character.encode("utf-8"))
            if length + size > 6000:
                chunks.append("".join(current))
                current, length = [], 0
            current.append(character)
            length += size
        chunks.append("".join(current))
        for number, chunk in enumerate(chunks, 1):
            yield unit(number, "Text" if len(chunks) == 1 else f"Part {number}", chunk)
    else:
        raise ValueError("This file type is not supported in V1.")
